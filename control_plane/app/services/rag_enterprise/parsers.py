import csv
import io
import logging
import os
from typing import List, Optional, Tuple

from app.services.rag_enterprise.schemas import ParseResult, ParserStatus, SUPPORTED_EXTENSIONS

logger = logging.getLogger(__name__)

_parser_availability = {}


def _check_dependency(name: str, package: str) -> bool:
    try:
        __import__(package)
        _parser_availability[name] = True
        return True
    except ImportError:
        _parser_availability[name] = False
        return False


_check_dependency("pypdf", "fitz")
_check_dependency("python-docx", "docx")
_check_dependency("openpyxl", "openpyxl")


def get_parser_status(ext: str) -> ParserStatus:
    status_map = {
        ".txt": ParserStatus(
            extension=".txt", supported=True, available=True,
            dependency=None, remediation=None,
        ),
        ".md": ParserStatus(
            extension=".md", supported=True, available=True,
            dependency=None, remediation=None,
        ),
        ".csv": ParserStatus(
            extension=".csv", supported=True, available=True,
            dependency=None, remediation=None,
        ),
        ".pdf": ParserStatus(
            extension=".pdf", supported=True,
            available=_parser_availability.get("pypdf", False),
            dependency="pymupdf (fitz)",
            remediation="pip install pymupdf",
        ),
        ".docx": ParserStatus(
            extension=".docx", supported=True,
            available=_parser_availability.get("python-docx", False),
            dependency="python-docx",
            remediation="pip install python-docx",
        ),
        ".xlsx": ParserStatus(
            extension=".xlsx", supported=True,
            available=_parser_availability.get("openpyxl", False),
            dependency="openpyxl",
            remediation="pip install openpyxl",
        ),
    }
    return status_map.get(ext, ParserStatus(
        extension=ext, supported=False, available=False,
        dependency=None, remediation="Unsupported file type",
    ))


async def parse_txt(file_path: str) -> ParseResult:
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        text = f.read()
    pages = [1] if text.strip() else []
    return ParseResult(text=text, pages=pages, metadata={"source": file_path})


async def parse_md(file_path: str) -> ParseResult:
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        text = f.read()
    pages = [1] if text.strip() else []
    return ParseResult(text=text, pages=pages, metadata={"source": file_path, "format": "markdown"})


async def parse_csv(file_path: str) -> ParseResult:
    pages = []
    all_text_parts = []
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        reader = csv.reader(f)
        for row_num, row in enumerate(reader, start=1):
            line = ", ".join(row)
            all_text_parts.append(line)
            pages.append(row_num)
    text = "\n".join(all_text_parts)
    return ParseResult(text=text, pages=pages, metadata={"source": file_path, "format": "csv"})


async def parse_pdf(file_path: str) -> ParseResult:
    if not _parser_availability.get("pypdf", False):
        raise ImportError(
            "PDF parsing requires pymupdf. Install with: pip install pymupdf"
        )
    import fitz

    text_by_page = []
    pages = []
    with fitz.open(file_path) as pdf:
        for page_num, page in enumerate(pdf, start=1):
            text = page.get_text().strip()
            if text:
                text_by_page.append(text)
                pages.append(page_num)
    text = "\n\n".join(text_by_page)
    return ParseResult(text=text, pages=pages, metadata={
        "source": file_path, "format": "pdf", "page_count": len(pdf),
    })


async def parse_docx(file_path: str) -> ParseResult:
    if not _parser_availability.get("python-docx", False):
        raise ImportError(
            "DOCX parsing requires python-docx. Install with: pip install python-docx"
        )
    from docx import Document

    doc = Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    text = "\n".join(paragraphs)
    pages = [1] if text.strip() else []
    return ParseResult(text=text, pages=pages, metadata={
        "source": file_path, "format": "docx",
    })


async def parse_xlsx(file_path: str) -> ParseResult:
    if not _parser_availability.get("openpyxl", False):
        raise ImportError(
            "XLSX parsing requires openpyxl. Install with: pip install openpyxl"
        )
    from openpyxl import load_workbook

    wb = load_workbook(file_path, read_only=True, data_only=True)
    all_text = []
    pages = []
    sheet_index = 0
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        sheet_text = []
        sheet_index += 1
        for row in ws.iter_rows(values_only=True):
            row_text = " | ".join(str(cell) for cell in row if cell is not None)
            if row_text.strip():
                sheet_text.append(row_text)
        if sheet_text:
            all_text.append(f"[Sheet: {sheet_name}]")
            all_text.extend(sheet_text)
            pages.append(sheet_index)
    text = "\n".join(all_text)
    wb.close()
    return ParseResult(text=text, pages=pages, metadata={
        "source": file_path, "format": "xlsx", "sheets": wb.sheetnames,
    })


PARSER_MAP = {
    ".txt": parse_txt,
    ".md": parse_md,
    ".csv": parse_csv,
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".xlsx": parse_xlsx,
}


async def parse_file(file_path: str, ext: str) -> ParseResult:
    ext = ext.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file extension: {ext}")

    parser = PARSER_MAP.get(ext)
    if parser is None:
        raise ValueError(f"No parser available for: {ext}")

    return await parser(file_path)


def get_available_parsers() -> List[ParserStatus]:
    return [get_parser_status(ext) for ext in sorted(SUPPORTED_EXTENSIONS)]


def get_parsers_summary() -> Tuple[List[str], List[str]]:
    available = []
    skipped = []
    for ext in sorted(SUPPORTED_EXTENSIONS):
        status = get_parser_status(ext)
        if status.available:
            available.append(ext)
        else:
            skipped.append(ext)
    return available, skipped
