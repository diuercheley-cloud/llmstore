#!/usr/bin/env python3
"""Validate and execute ad hoc test scripts stored under scripts/."""

from __future__ import annotations

import argparse
import datetime
import json
import os
import subprocess
import sys
import time
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Iterable

try:
    from docx import Document
except ModuleNotFoundError:  # pragma: no cover - optional dependency
    Document = None


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS_DIR = ROOT / "scripts"
REPORT_DIR = ROOT / "artifacts" / "scripts" / "latest"
DEFAULT_TIMEOUT_S = 120
TEST_PATTERNS = ("test_*.py", "test-*.py", "test_*.sh", "test-*.sh")


@dataclass
class CheckResult:
    name: str
    kind: str
    status: str
    checks: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    exit_code: int | None = None
    duration_s: float | None = None
    command: list[str] = field(default_factory=list)


def discover_script_tests() -> list[Path]:
    discovered: set[Path] = set()
    for pattern in TEST_PATTERNS:
        discovered.update(path for path in SCRIPTS_DIR.glob(pattern) if path.is_file())
    return sorted(discovered)


def has_shebang(path: Path) -> bool:
    try:
        first_line = path.read_text(encoding="utf-8", errors="ignore").splitlines()[0]
    except (IndexError, OSError):
        return False
    return first_line.startswith("#!")


def run_command(cmd: list[str], timeout_s: int) -> tuple[int, str]:
    completed = subprocess.run(
        cmd,
        cwd=ROOT,
        env=os.environ.copy(),
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=timeout_s,
        check=False,
    )
    return completed.returncode, completed.stdout


def static_validate(path: Path) -> CheckResult:
    result = CheckResult(name=str(path.relative_to(ROOT)), kind=path.suffix.lstrip("."), status="pass")

    if path.suffix == ".sh":
        if not os.access(path, os.X_OK):
            result.status = "fail"
            result.checks.append("shell script is not executable")
        else:
            result.checks.append("executable bit present")

        if has_shebang(path):
            result.checks.append("shebang present")
        else:
            result.status = "fail"
            result.checks.append("missing shebang")

        code, output = run_command(["bash", "-n", str(path)], timeout_s=30)
        if code == 0:
            result.checks.append("bash -n passed")
        else:
            result.status = "fail"
            result.checks.append("bash -n failed")
            stripped = output.strip()
            if stripped:
                result.notes.append(stripped)

        shellcheck = shutil_which("shellcheck")
        if shellcheck:
            code, output = run_command([shellcheck, str(path)], timeout_s=30)
            if code == 0:
                result.checks.append("shellcheck passed")
            else:
                result.status = "fail"
                result.checks.append("shellcheck failed")
                stripped = output.strip()
                if stripped:
                    result.notes.append(stripped)
        else:
            result.checks.append("shellcheck not available")

    elif path.suffix == ".py":
        code, output = run_command([sys.executable, "-m", "py_compile", str(path)], timeout_s=30)
        if code == 0:
            result.checks.append("py_compile passed")
        else:
            result.status = "fail"
            result.checks.append("py_compile failed")
            stripped = output.strip()
            if stripped:
                result.notes.append(stripped)
    else:
        result.status = "skip"
        result.checks.append("unsupported file type")

    return result


def execute_script(path: Path, timeout_s: int, verbose: bool) -> CheckResult:
    result = CheckResult(name=str(path.relative_to(ROOT)), kind=path.suffix.lstrip("."), status="pass")
    start = time.perf_counter()

    if path.suffix == ".sh":
        cmd = ["bash", str(path)]
    elif path.suffix == ".py":
        cmd = [sys.executable, "-u", str(path)]
    else:
        result.status = "skip"
        result.checks.append("unsupported file type")
        return result

    result.command = cmd
    if verbose:
        print(f"\n[RUN] {result.name}")
        print(f"      {' '.join(cmd)}")

    try:
        code, output = run_command(cmd, timeout_s=timeout_s)
        result.exit_code = code
        if output:
            if verbose:
                print(output, end="" if output.endswith("\n") else "\n")
        if code == 0:
            result.checks.append("execution passed")
        else:
            result.status = "fail"
            result.checks.append("execution failed")
            stripped = output.strip()
            if stripped:
                result.notes.append(stripped)
    except subprocess.TimeoutExpired as exc:
        result.status = "fail"
        result.exit_code = 124
        result.checks.append(f"timed out after {timeout_s}s")
        output = "".join(part for part in (exc.stdout, exc.stderr) if part)
        if output:
            if verbose:
                print(output, end="" if output.endswith("\n") else "\n")
            result.notes.append(output.strip())
    finally:
        result.duration_s = round(time.perf_counter() - start, 3)

    if verbose:
        status_label = "PASSED" if result.status == "pass" else "FAILED"
        print(f"[{status_label}] {result.name} ({result.duration_s:.3f}s)")

    return result


def run_script_test(path: Path, timeout_s: int, verbose: bool) -> CheckResult:
    static_result = static_validate(path)

    if static_result.status == "skip":
        return static_result

    if verbose:
        print(f"\n=== {static_result.name} ===")
        print(f"Static checks: {', '.join(static_result.checks) if static_result.checks else 'none'}")

    execution_result = execute_script(path, timeout_s=timeout_s, verbose=verbose)

    combined = CheckResult(
        name=static_result.name,
        kind=static_result.kind,
        status="pass",
        checks=static_result.checks + execution_result.checks,
        notes=static_result.notes + execution_result.notes,
        exit_code=execution_result.exit_code,
        duration_s=execution_result.duration_s,
        command=execution_result.command,
    )
    if static_result.status == "fail" or execution_result.status == "fail":
        combined.status = "fail"

    if verbose and static_result.notes:
        print("Static notes:")
        for note in static_result.notes:
            print(note)

    return combined


def shutil_which(command: str) -> str | None:
    from shutil import which

    return which(command)


def write_report(results: list[CheckResult], execute_mode: bool, timeout_s: int) -> Path:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    report_path = REPORT_DIR / "script-tests-validation.md"
    json_path = REPORT_DIR / "script-tests-validation.json"

    totals = {
        "total": len(results),
        "passed": sum(1 for result in results if result.status == "pass"),
        "failed": sum(1 for result in results if result.status == "fail"),
        "skipped": sum(1 for result in results if result.status == "skip"),
    }

    lines: list[str] = [
        "# Script Tests Validation Report",
        "",
        f"- Mode: {'execute' if execute_mode else 'static'}",
        f"- Timeout: {timeout_s}s" if execute_mode else "- Timeout: n/a",
        f"- Total files: {totals['total']}",
        f"- Passed: {totals['passed']}",
        f"- Failed: {totals['failed']}",
        f"- Skipped: {totals['skipped']}",
        "",
        "| File | Kind | Status | Checks | Notes |",
        "| --- | --- | --- | --- | --- |",
    ]

    for result in results:
        checks = "<br>".join(result.checks) if result.checks else "-"
        notes = "<br>".join(result.notes) if result.notes else "-"
        lines.append(f"| `{result.name}` | `{result.kind}` | `{result.status}` | {checks} | {notes} |")

    report_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    json_path.write_text(
        json.dumps(
            {
                "mode": "execute" if execute_mode else "static",
                "timeout_s": timeout_s if execute_mode else None,
                "totals": totals,
                "results": [asdict(result) for result in results],
            },
            indent=2,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )
    return report_path


def generate_docx_report(results: list[CheckResult], execute_mode: bool, timeout_s: int) -> Path | None:
    if Document is None:
        print("python-docx is not available; skipping .docx report generation.")
        return None

    doc = Document()
    doc.add_heading("Relatorio de Execucao de Testes", 0)
    doc.add_paragraph(f"Data da execucao: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph(f"Modo: {'execute' if execute_mode else 'static'}")
    doc.add_paragraph(f"Timeout: {timeout_s}s" if execute_mode else "Timeout: n/a")

    totals = {
        "total": len(results),
        "passed": sum(1 for result in results if result.status == "pass"),
        "failed": sum(1 for result in results if result.status == "fail"),
        "skipped": sum(1 for result in results if result.status == "skip"),
    }

    doc.add_paragraph(f"Total: {totals['total']}")
    doc.add_paragraph(f"Passed: {totals['passed']}")
    doc.add_paragraph(f"Failed: {totals['failed']}")
    doc.add_paragraph(f"Skipped: {totals['skipped']}")

    for result in results:
        heading = f"{result.name} - {result.status.upper()}"
        doc.add_heading(heading, level=1)
        doc.add_paragraph(f"Kind: {result.kind}")
        if result.command:
            doc.add_paragraph(f"Command: {' '.join(result.command)}")
        if result.exit_code is not None:
            doc.add_paragraph(f"Exit code: {result.exit_code}")
        if result.duration_s is not None:
            doc.add_paragraph(f"Duration: {result.duration_s:.3f}s")
        if result.checks:
            doc.add_paragraph("Checks: " + ", ".join(result.checks))
        if result.notes:
            doc.add_paragraph("Notes:")
            for note in result.notes[-5:]:
                doc.add_paragraph(note)

    report_name = f"relatorio_testes_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    doc.save(report_name)
    return Path(report_name)


def print_summary(results: list[CheckResult], execute_mode: bool, report_path: Path) -> int:
    passed = sum(1 for result in results if result.status == "pass")
    failed = sum(1 for result in results if result.status == "fail")
    skipped = sum(1 for result in results if result.status == "skip")
    mode = "execute" if execute_mode else "static"

    print("\nSummary")
    print(f"Mode: {mode}")
    print(f"Validated files: {len(results)}")
    print(f"Passed: {passed}")
    print(f"Failed: {failed}")
    print(f"Skipped: {skipped}")
    print(f"Report: {report_path}")

    for result in results:
        print(f"- {result.name}: {result.status.upper()}")

    if failed:
        print("\nFailures:")
        for result in results:
            if result.status != "fail":
                continue
            print(f"- {result.name}")
            for note in result.notes[:3]:
                if note:
                    print(f"  {note}")

    return 0 if failed == 0 else 1


def main(argv: Iterable[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Validate and run the test scripts stored under scripts/.")
    parser.add_argument(
        "--static-only",
        action="store_true",
        help="run only static checks without executing the scripts",
    )
    parser.add_argument(
        "--execute",
        action="store_true",
        help="explicitly run scripts after static checks (execution is the default)",
    )
    parser.add_argument(
        "--timeout",
        type=int,
        default=DEFAULT_TIMEOUT_S,
        help=f"timeout per script in execute mode (default: {DEFAULT_TIMEOUT_S}s)",
    )
    parser.add_argument(
        "--quiet",
        action="store_true",
        help="suppress per-script command output",
    )
    args = parser.parse_args(list(argv) if argv is not None else None)

    execute_mode = not args.static_only
    if args.execute:
        execute_mode = True

    script_tests = discover_script_tests()
    if not script_tests:
        print("No test scripts found under scripts/.")
        return 0

    print(f"Found {len(script_tests)} test script(s) under scripts/.")

    results: list[CheckResult] = []
    for path in script_tests:
        static_result = static_validate(path)
        if not execute_mode:
            results.append(static_result)
            print(f"- {static_result.name}: {static_result.status.upper()}")
            continue

        result = run_script_test(path, timeout_s=args.timeout, verbose=not args.quiet)
        results.append(result)

    report_path = write_report(results, execute_mode=execute_mode, timeout_s=args.timeout)
    generate_docx_report(results, execute_mode=execute_mode, timeout_s=args.timeout)
    return print_summary(results, execute_mode=execute_mode, report_path=report_path)


if __name__ == "__main__":
    raise SystemExit(main())
